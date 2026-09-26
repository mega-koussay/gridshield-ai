"""Render PDF deliverables with ReportLab (exec summary, report, deck).

Light markdown parsing: headings, bullets, tables, bold — enough for the
documents in deliverables/, keeping everything reproducible from source.
"""
from __future__ import annotations

import re
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parents[1]
DELIV = ROOT / "deliverables"

INK = colors.HexColor("#0b1220")
SLATE = colors.HexColor("#334155")
BLUE = colors.HexColor("#0284c7")


def _esc(t: str) -> str:
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"\*(.+?)\*", r"<i>\1</i>", t)
    return t


def md_to_pdf(md_path: Path, pdf_path: Path) -> None:
    styles = getSampleStyleSheet()
    h0 = ParagraphStyle("H0", parent=styles["Title"], fontSize=20, leading=24,
                        textColor=INK, spaceAfter=8)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=14, leading=18,
                        textColor=BLUE, spaceBefore=10, spaceAfter=4)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=11.5, leading=15,
                        textColor=INK, spaceBefore=8, spaceAfter=3)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.3, leading=12.6,
                          textColor=colors.HexColor("#111827"))
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=12, bulletIndent=4,
                            spaceAfter=2)
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.6, leading=11)

    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(SLATE)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(18 * mm, 12 * mm, "GridShield AI — IEEE SmartSecureGrid Challenge 2026 PoC")
        canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, f"{canvas.getPageNumber()}")
        canvas.restoreState()

    doc = BaseDocTemplate(str(pdf_path), pagesize=A4,
                          leftMargin=20 * mm, rightMargin=20 * mm,
                          topMargin=16 * mm, bottomMargin=18 * mm,
                          title=md_path.stem)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([PageTemplate(id="pt", frames=[frame], onPage=header_footer)])

    story = []
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            story.append(Paragraph(_esc(line), ParagraphStyle(
                "Code", parent=body, fontName="Courier", fontSize=7.6, leading=9.4)))
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            story.append(Paragraph(_esc(line[2:]), h0))
        elif line.startswith("## "):
            story.append(Paragraph(_esc(line[3:]), h1))
        elif line.startswith("### "):
            story.append(Paragraph(_esc(line[4:]), h2))
        elif line.strip() in ("---", "***"):
            story.append(Spacer(1, 4))
        elif line.startswith("|"):
            # collect table block
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                data = [[Paragraph(_esc(c), cell) for c in r] for r in rows]
                t = Table(data, colWidths="*")
                t.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, SLATE),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0f2fe")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ]))
                story.append(t)
                story.append(Spacer(1, 6))
            continue
        elif line.startswith("- "):
            story.append(Paragraph(_esc(line[2:]), bullet, bulletText="•"))
        elif re.match(r"^\d+\. ", line):
            num = re.match(r"^(\d+)\. ", line).group(1)
            story.append(Paragraph(_esc(re.sub(r"^\d+\. ", "", line)), bullet,
                                   bulletText=num + "."))
        elif line.startswith(">"):
            story.append(Paragraph(_esc(line.lstrip("> ")), ParagraphStyle(
                "Q", parent=body, textColor=SLATE, leftIndent=10)))
        else:
            story.append(Paragraph(_esc(line.strip()), body))
        i += 1

    doc.build(story)
    print("wrote", pdf_path)


def deck_to_pdf(pptx_path: Path, pdf_path: Path) -> None:
    """Render each slide's text content as a PDF page (print-friendly)."""
    from pptx import Presentation
    sys_style = getSampleStyleSheet()
    title = ParagraphStyle("T", parent=sys_style["Title"], fontSize=30, textColor=INK)
    sub = ParagraphStyle("S", parent=sys_style["Heading2"], fontSize=15,
                         textColor=BLUE, spaceAfter=14)
    item = ParagraphStyle("I", parent=sys_style["BodyText"], fontSize=12.5,
                          leading=19, textColor=colors.HexColor("#111827"))

    def hf(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(INK)
        canvas.rect(0, 0, A4[0], A4[1], stroke=0, fill=0) if False else None
        canvas.setFillColor(SLATE)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 12 * mm, "GridShield AI — Pitch Deck")
        canvas.restoreState()

    doc = BaseDocTemplate(str(pdf_path), pagesize=A4,
                          leftMargin=22 * mm, rightMargin=22 * mm,
                          topMargin=20 * mm, bottomMargin=20 * mm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([PageTemplate(id="pt", frames=[frame], onPage=hf)])

    prs = Presentation(str(pptx_path))
    story = []
    for idx, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs).strip()
                    if txt:
                        texts.append(txt)
        if not texts:
            continue
        kicker = texts[0]
        t = texts[1] if len(texts) > 1 else ""
        s = texts[2] if len(texts) > 2 else ""
        story.append(Paragraph(_esc(kicker), ParagraphStyle(
            "K", parent=item, fontSize=9, textColor=SLATE)))
        story.append(Spacer(1, 6))
        story.append(Paragraph(_esc(t), title))
        story.append(Spacer(1, 6))
        if s:
            story.append(Paragraph(_esc(s), sub))
        for b in texts[3:]:
            story.append(Paragraph("▸  " + _esc(b), item))
            story.append(Spacer(1, 4))
        if idx < len(prs.slides.__iter__.__self__._sldIdLst) - 1:  # page break between slides
            from reportlab.platypus import PageBreak
            story.append(PageBreak())
    doc.build(story)
    print("wrote", pdf_path)


if __name__ == "__main__":
    md_to_pdf(DELIV / "executive_summary.md", DELIV / "executive_summary.pdf")
    md_to_pdf(DELIV / "technical_report.md", DELIV / "technical_report.pdf")
    deck_to_pdf(DELIV / "pitch_deck.pptx", DELIV / "pitch_deck.pdf")
